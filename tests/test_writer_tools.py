import json
from pathlib import Path
from unittest import mock

from docx import Document as DocxDocument

from mcp_documents_reader import (
    write_presentation,
    write_spreadsheet,
    write_word_document,
)


def test_write_word_document_docx_with_tables(temp_document_dir: str) -> None:
    output_path = Path(temp_document_dir) / "generated.docx"

    result = write_word_document(
        filename=str(output_path),
        title="周报",
        paragraphs=["第一段", "第二段"],
        tables=[
            {
                "title": "数据表",
                "headers": ["姓名", "年龄"],
                "rows": [["张三", 18], ["李四", 20]],
            }
        ],
    )

    payload = json.loads(result)
    assert payload["format"] == "docx"
    assert Path(payload["path"]).exists()

    document = DocxDocument(str(output_path))
    paragraph_texts = [paragraph.text for paragraph in document.paragraphs]
    assert "周报" in paragraph_texts
    assert "第一段" in paragraph_texts
    assert "数据表" in paragraph_texts
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "姓名"
    assert document.tables[0].cell(1, 0).text == "张三"


@mock.patch("mcp_documents_reader._convert_with_libreoffice")
def test_write_word_document_doc_uses_conversion(
    mock_convert: mock.MagicMock, temp_document_dir: str
) -> None:
    output_path = Path(temp_document_dir) / "generated.doc"
    output_path.write_bytes(b"fake-doc")
    mock_convert.return_value = output_path

    result = write_word_document(
        filename=str(output_path),
        title="文档",
        paragraphs=["内容"],
    )

    payload = json.loads(result)
    assert payload["format"] == "doc"
    assert payload["source_format"] == "docx"
    mock_convert.assert_called_once()


@mock.patch("mcp_documents_reader._convert_with_libreoffice", return_value=None)
def test_write_word_document_doc_conversion_unavailable(
    mock_convert: mock.MagicMock, temp_document_dir: str
) -> None:
    output_path = Path(temp_document_dir) / "generated.doc"
    result = write_word_document(filename=str(output_path), title="文档")

    assert "LibreOffice conversion is unavailable" in result
    mock_convert.assert_called_once()


def test_write_word_document_invalid_extension(temp_document_dir: str) -> None:
    output_path = Path(temp_document_dir) / "generated.txt"
    result = write_word_document(filename=str(output_path), title="文档")

    assert result == "Error: Word generation supports .docx and .doc output only."


@mock.patch("mcp_documents_reader._load_presentation_dependencies")
def test_write_presentation_pptx(
    mock_load_dependencies: mock.MagicMock, temp_document_dir: str
) -> None:
    output_path = Path(temp_document_dir) / "generated.pptx"

    class FakeParagraph:
        def __init__(self) -> None:
            self.text = ""
            self.level = None

    class FakeTextFrame:
        def __init__(self) -> None:
            self.text = ""
            self.paragraphs = [FakeParagraph()]

        def add_paragraph(self) -> FakeParagraph:
            paragraph = FakeParagraph()
            self.paragraphs.append(paragraph)
            return paragraph

    class FakeTextBox:
        def __init__(self) -> None:
            self.text_frame = FakeTextFrame()

    class FakeCell:
        def __init__(self) -> None:
            self.text = ""

    class FakeTable:
        def __init__(self, rows: int, cols: int) -> None:
            self._cells = [[FakeCell() for _ in range(cols)] for _ in range(rows)]

        def cell(self, row: int, col: int) -> FakeCell:
            return self._cells[row][col]

    class FakeTableShape:
        def __init__(self, rows: int, cols: int) -> None:
            self.table = FakeTable(rows, cols)

    class FakeTitleShape:
        def __init__(self) -> None:
            self.has_text_frame = True
            self.text_frame = FakeTextFrame()

    class FakeSlideShapes:
        def __init__(self, has_title: bool = False) -> None:
            self.title = FakeTitleShape() if has_title else None
            self.textboxes: list[FakeTextBox] = []
            self.tables: list[FakeTableShape] = []

        def add_textbox(self, *args, **kwargs) -> FakeTextBox:
            textbox = FakeTextBox()
            self.textboxes.append(textbox)
            return textbox

        def add_table(self, rows: int, cols: int, *args, **kwargs) -> FakeTableShape:
            table_shape = FakeTableShape(rows, cols)
            self.tables.append(table_shape)
            return table_shape

    class FakePlaceholder:
        def __init__(self) -> None:
            self.has_text_frame = True
            self.text_frame = FakeTextFrame()

    class FakeSlide:
        def __init__(self, has_title: bool = False, placeholder_count: int = 0) -> None:
            self.shapes = FakeSlideShapes(has_title=has_title)
            self.placeholders = [FakePlaceholder() for _ in range(placeholder_count)]

    class FakeSlides(list):
        def __init__(self) -> None:
            super().__init__()
            self.created_slides: list[FakeSlide] = []

        def add_slide(self, layout: object) -> FakeSlide:
            if layout == "title-layout":
                slide = FakeSlide(has_title=True, placeholder_count=2)
            else:
                slide = FakeSlide()
            self.created_slides.append(slide)
            self.append(slide)
            return slide

    class FakePresentation:
        def __init__(self) -> None:
            self.slide_layouts = ["title-layout"] + ["content-layout"] * 6
            self.slides = FakeSlides()
            self.saved_path: str | None = None

        def save(self, path: str) -> None:
            self.saved_path = path
            Path(path).write_bytes(b"fake-pptx")

    fake_presentation = FakePresentation()
    mock_load_dependencies.return_value = (
        lambda: fake_presentation,
        lambda value: value,
    )

    result = write_presentation(
        filename=str(output_path),
        title="季度汇报",
        subtitle="Q2",
        slides=[
            {
                "title": "第一页",
                "paragraphs": ["概述"],
                "bullets": ["重点1", "重点2"],
                "table": {
                    "headers": ["项目", "进度"],
                    "rows": [["需求", "完成"]],
                },
            }
        ],
    )

    payload = json.loads(result)
    assert payload["format"] == "pptx"
    assert Path(payload["path"]).exists()
    assert fake_presentation.saved_path == str(output_path)
    assert len(fake_presentation.slides.created_slides) == 2

    title_slide = fake_presentation.slides.created_slides[0]
    assert title_slide.shapes.title is not None
    assert title_slide.shapes.title.text_frame.text == "季度汇报"
    assert title_slide.placeholders[1].text_frame.text == "Q2"

    content_slide = fake_presentation.slides.created_slides[1]
    assert content_slide.shapes.textboxes[0].text_frame.text == "第一页"
    paragraphs = content_slide.shapes.textboxes[1].text_frame.paragraphs
    assert paragraphs[0].text == "概述"
    assert paragraphs[1].text == "重点1"
    assert paragraphs[2].text == "重点2"
    table = content_slide.shapes.tables[0].table
    assert table.cell(0, 0).text == "项目"
    assert table.cell(1, 1).text == "完成"


@mock.patch("mcp_documents_reader._load_presentation_dependencies")
@mock.patch("mcp_documents_reader._convert_with_libreoffice")
def test_write_presentation_ppt_uses_conversion(
    mock_convert: mock.MagicMock,
    mock_load_dependencies: mock.MagicMock,
    temp_document_dir: str,
) -> None:
    output_path = Path(temp_document_dir) / "generated.ppt"
    output_path.write_bytes(b"fake-ppt")

    class FakeSlides(list):
        def add_slide(self, layout: object):
            slide = mock.MagicMock()
            slide.shapes.title = None
            slide.placeholders = []
            self.append(slide)
            return slide

    class FakePresentation:
        def __init__(self) -> None:
            self.slide_layouts = [object()] * 7
            self.slides = FakeSlides()

        def save(self, path: str) -> None:
            Path(path).write_bytes(b"fake-pptx")

    mock_load_dependencies.return_value = (
        lambda: FakePresentation(),
        lambda value: value,
    )
    mock_convert.return_value = output_path

    result = write_presentation(filename=str(output_path), title="演示")

    payload = json.loads(result)
    assert payload["format"] == "ppt"
    assert payload["source_format"] == "pptx"
    mock_convert.assert_called_once()


@mock.patch(
    "mcp_documents_reader._load_presentation_dependencies",
    side_effect=RuntimeError("missing"),
)
def test_write_presentation_dependency_error(
    mock_load_dependencies: mock.MagicMock, temp_document_dir: str
) -> None:
    output_path = Path(temp_document_dir) / "generated.pptx"
    result = write_presentation(filename=str(output_path), title="演示")

    assert result == "Error writing presentation: missing"
    mock_load_dependencies.assert_called_once()


def test_write_presentation_invalid_extension(temp_document_dir: str) -> None:
    output_path = Path(temp_document_dir) / "generated.pdf"
    result = write_presentation(filename=str(output_path), title="演示")

    assert (
        result == "Error: Presentation generation supports .pptx and .ppt output only."
    )


def test_write_spreadsheet_xlsx(temp_document_dir: str) -> None:
    output_path = Path(temp_document_dir) / "data.xlsx"

    result = write_spreadsheet(
        filename=str(output_path),
        sheets=[
            {
                "name": "员工",
                "headers": ["姓名", "年龄"],
                "rows": [["张三", 18], ["李四", 20]],
            },
            {
                "name": "产品",
                "headers": ["名称", "价格"],
                "rows": [["苹果", 3.5]],
            },
        ],
    )

    payload = json.loads(result)
    assert payload["format"] == "xlsx"
    assert Path(payload["path"]).exists()

    from openpyxl import load_workbook

    workbook = load_workbook(str(output_path))
    assert workbook.sheetnames == ["员工", "产品"]
    sheet = workbook["员工"]
    assert sheet.cell(1, 1).value == "姓名"
    assert sheet.cell(2, 1).value == "张三"
    assert sheet.cell(2, 2).value == 18


def test_write_spreadsheet_shorthand_single_sheet(temp_document_dir: str) -> None:
    output_path = Path(temp_document_dir) / "simple.xlsx"

    result = write_spreadsheet(
        filename=str(output_path),
        headers=["名称"],
        rows=[["测试"]],
    )

    payload = json.loads(result)
    assert payload["format"] == "xlsx"

    from openpyxl import load_workbook

    workbook = load_workbook(str(output_path))
    sheet = workbook.active
    assert sheet is not None
    assert sheet.cell(1, 1).value == "名称"
    assert sheet.cell(2, 1).value == "测试"


def test_write_spreadsheet_csv(temp_document_dir: str) -> None:
    output_path = Path(temp_document_dir) / "data.csv"

    result = write_spreadsheet(
        filename=str(output_path),
        headers=["姓名", "年龄"],
        rows=[["张三", 18]],
    )

    payload = json.loads(result)
    assert payload["format"] == "csv"
    content = output_path.read_text(encoding="utf-8-sig")
    assert "姓名,年龄" in content
    assert "张三,18" in content


def test_write_spreadsheet_csv_multiple_sheets_error(temp_document_dir: str) -> None:
    output_path = Path(temp_document_dir) / "data.csv"

    result = write_spreadsheet(
        filename=str(output_path),
        sheets=[{"name": "a"}, {"name": "b"}],
    )

    assert "single sheet" in result


def test_write_spreadsheet_invalid_extension(temp_document_dir: str) -> None:
    output_path = Path(temp_document_dir) / "data.txt"

    result = write_spreadsheet(filename=str(output_path), rows=[["a"]])

    assert (
        "Error: Spreadsheet generation supports .xlsx, .csv and .xls output only."
        in result
    )


@mock.patch("mcp_documents_reader._convert_with_libreoffice")
def test_write_spreadsheet_xls_uses_conversion(
    mock_convert: mock.MagicMock, temp_document_dir: str
) -> None:
    output_path = Path(temp_document_dir) / "data.xls"
    output_path.write_bytes(b"fake-xls")
    mock_convert.return_value = output_path

    result = write_spreadsheet(filename=str(output_path), rows=[["a", 1]])

    payload = json.loads(result)
    assert payload["format"] == "xls"
    assert payload["source_format"] == "xlsx"
    mock_convert.assert_called_once()


@mock.patch("mcp_documents_reader._convert_with_libreoffice", return_value=None)
def test_write_spreadsheet_xls_conversion_unavailable(
    mock_convert: mock.MagicMock, temp_document_dir: str
) -> None:
    output_path = Path(temp_document_dir) / "data.xls"

    result = write_spreadsheet(filename=str(output_path), rows=[["a"]])

    assert "LibreOffice conversion is unavailable" in result
    mock_convert.assert_called_once()


def test_write_spreadsheet_sanitizes_sheet_name(temp_document_dir: str) -> None:
    output_path = Path(temp_document_dir) / "names.xlsx"

    result = write_spreadsheet(
        filename=str(output_path),
        sheets=[{"name": "非法[名]称:*?", "rows": [["x"]]}],
    )

    payload = json.loads(result)
    assert payload["format"] == "xlsx"

    from openpyxl import load_workbook

    workbook = load_workbook(str(output_path))
    assert workbook.sheetnames == ["非法名称"]
